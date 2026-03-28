"""Generate docs/Fitness_Band_Data_Format.docx — Tanish: nature of data (Q2, first part)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    out = root / "docs" / "Fitness_Band_Data_Format.docx"

    doc = Document()
    doc.core_properties.author = "Tanish"
    doc.core_properties.title = "What the Fitness Band Data Looks Like"
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("What the Fitness Band Data Looks Like", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        "FitFuel — CMPE 195A/B Senior Design | San José State University | Spring 2026"
    )
    run.italic = True

    auth = doc.add_paragraph()
    auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    auth.add_run("Tanish").bold = True

    doc.add_paragraph()

    q = doc.add_paragraph()
    q.add_run("Question. ").bold = True
    q.add_run(
        "What will the data look like—data from the sensor, raw data, or computed summary, "
        "and in what format?"
    )

    doc.add_heading("Answer", level=1)
    doc.add_paragraph(
        "At the sensor level, the fitness band produces raw measurements: for example, optical "
        "PPG samples for heart rate and accelerometer and gyroscope data for motion. The "
        "microcontroller reads these registers and processes them. The design does not send "
        "every raw sample over the radio link; doing so would consume excessive power and "
        "bandwidth relative to the benefit for our use case."
    )
    doc.add_paragraph(
        "Instead, over Bluetooth Low Energy (BLE), the band transmits a "
    )
    p = doc.add_paragraph()
    p.add_run("compact computed summary").bold = True
    p.add_run(
        ": aggregated values such as heart rate in beats per minute, step count, and an "
        "estimate of calories burned, together with supporting fields (for example, battery "
        "level and a sequence number for integrity). That summary reflects processed sensor "
        "data rather than a raw sensor dump."
    )
    doc.add_paragraph(
        "The specific encoding, field sizes, and byte order used on the wireless link for "
        "host software are documented separately by Ishan Sikka in "
        "Telemetry_Data_Format_Ishan_Sikka.docx."
    )

    doc.add_paragraph()
    ref = doc.add_paragraph()
    ref.add_run("Reference. ").bold = True
    ref.add_run("docs/hardware-fitness-band-connectivity.md")

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()

"""Generate docs/Telemetry_Data_Format_Ishan_Sikka.docx — format specification (Q2, second part)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    out = root / "docs" / "Telemetry_Data_Format_Ishan_Sikka.docx"

    doc = Document()
    doc.core_properties.author = "Ishan Sikka"
    doc.core_properties.title = "Fitness Band Telemetry: Data Format Specification"
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading(
        "Fitness Band Telemetry: Data Format Specification",
        level=0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        "FitFuel — CMPE 195A/B Senior Design | San José State University | Spring 2026"
    )
    run.italic = True

    auth = doc.add_paragraph()
    auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    auth.add_run("Ishan Sikka").bold = True

    doc.add_paragraph()

    doc.add_heading("Purpose", level=1)
    doc.add_paragraph(
        "This document specifies "
    )
    p1 = doc.add_paragraph()
    p1.add_run("in what format ").bold = True
    p1.add_run(
        "telemetry data are represented when transmitted from the fitness band to a host "
        "system (mobile device or personal computer). It complements the parallel summary "
        "on sensor versus summary data authored by Tanish: here the focus is exclusively "
        "on encoding, layout, and host-side interpretation—not on whether samples are raw or "
        "aggregated at the source."
    )

    doc.add_heading("Transport and encapsulation", level=1)
    t1 = doc.add_paragraph()
    t1.add_run("Telemetry is delivered over Bluetooth Low Energy using GATT ")
    t1.add_run("notifications").bold = True
    t1.add_run(" on a project-defined ")
    t1.add_run("Telemetry").bold = True
    t1.add_run(
        " characteristic within a custom service. The application payload is a "
    )
    t1.add_run("binary record").bold = True
    t1.add_run(
        ", not JSON or XML; compact fixed-width fields minimize airtime and parsing cost. "
        "Service and characteristic 128-bit UUIDs will be published in a shared constants "
        "module referenced by firmware and client software (see project connectivity "
        "documentation)."
    )

    doc.add_heading("Binary layout (application payload)", level=1)
    le = doc.add_paragraph()
    le.add_run("All multi-byte integers use ")
    le.add_run("little-endian").bold = True
    le.add_run(
        " byte order. The notification payload is structured as follows (approximately "
        "twelve octets)."
    )

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Byte offset"
    hdr[1].text = "Field"
    hdr[2].text = "Type"
    hdr[3].text = "Semantics"

    rows = [
        ("0", "Packet type", "uint8", "Message class (e.g., 0x01 = telemetry)"),
        ("1–2", "Heart rate", "uint16", "Beats per minute"),
        ("3–4", "Steps", "uint16", "Step count since reset"),
        ("5–8", "Calories", "uint32", "Energy as kilocalories × 10 (divide by 10 for kcal)"),
        ("9", "Battery", "uint8", "State of charge, 0–100%"),
        ("10–11", "Sequence", "uint16", "Monotonic counter; gaps imply missed notifications"),
    ]
    for r in rows:
        row = table.add_row().cells
        for i, t in enumerate(r):
            row[i].text = t

    doc.add_paragraph()
    doc.add_paragraph(
        "This record fits typical default ATT MTU limits; larger payloads are possible after "
        "MTU negotiation between Central and Peripheral."
    )

    doc.add_heading("Host software interpretation", level=1)
    doc.add_paragraph(
        "On Android, the GATT callback supplies the characteristic value as a byte array; "
        "fields are read using unsigned little-endian interpretation consistent with the table "
        "above. On desktop systems, an equivalent layout may be parsed with Python struct "
        "unpacking—for example, format string <B H H I B H in little-endian order (12 bytes "
        "total)—when using a BLE library such as bleak."
    )
    doc.add_paragraph(
        "Standard Bluetooth Battery Service characteristics, if exposed, follow Bluetooth SIG "
        "data types and are orthogonal to this custom telemetry record."
    )

    doc.add_heading("Reference", level=1)
    doc.add_paragraph("docs/hardware-fitness-band-connectivity.md")

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
